#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate a formatted .docx undergraduate thesis proposal (开题报告).

Formatting follows 人文社科类本科毕业论文（设计）的常见格式口径（缺省值，见 references/format_spec.md）。
若用户提供了自己学校的模板/规范，请以学校规范为准并相应修改本文件顶部的字号常量:
  body  : 宋体 / Times New Roman 小四(12pt), 两端对齐, 首行缩进2字符, 行距固定值20磅
  H1    : 黑体 三号(16pt) 居中, 段前段后各空1行, 序号「一、」
  H2    : 黑体 小三(15pt), 序号「（一）」
  H3    : 黑体 四号(14pt), 序号「1.」
  H4    : 黑体 小四(12pt), 序号「（1）」
  refs  : 五号(10.5pt), 悬挂缩进2字符
  cover : 题目 小二(18pt) 小标宋; 信息 小三(15pt) 楷体

Usage:
    python generate_docx.py --output 开题报告.docx --data content.json
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERROR: python-docx is required. Install with: pip install python-docx")
    sys.exit(1)


# ---------------------------------------------------------------- constants --
SONG = '宋体'
HEI = '黑体'
KAI = '楷体'
XBS = '方正小标宋简体'          # 小标宋；若本机未装，Word 会回退，需人工检查
EN = 'Times New Roman'

PT_XIAOSI = Pt(12)            # 小四
PT_WU = Pt(10.5)              # 五号
LINE_EXACT_20 = Pt(20)        # 行距固定值 20 磅

BODY_INDENT = Pt(24)          # 首行缩进 2 字符（小四 12pt × 2）
REF_HANGING = Pt(21)          # 悬挂缩进 2 字符（五号 10.5pt × 2）

CN_DIGITS = ['', '一', '二', '三', '四', '五', '六', '七', '八', '九', '十']


def cn_number(n):
    """1..20 -> 一..二十"""
    if 1 <= n <= 10:
        return CN_DIGITS[n]
    if 11 <= n <= 19:
        return '十' + CN_DIGITS[n - 10]
    if n == 20:
        return '二十'
    return str(n)


def set_run_font(run, cn_font=SONG, size=PT_XIAOSI, bold=False,
                 en_font=EN, color=None):
    run.font.size = size
    run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = en_font
    rpr = run._element.get_or_add_rPr()
    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:eastAsia'), cn_font)
    rfonts.set(qn('w:ascii'), en_font)
    rfonts.set(qn('w:hAnsi'), en_font)
    rpr.insert(0, rfonts)


class DocBuilder:
    def __init__(self):
        self.doc = Document()
        self._setup_page()
        self._setup_normal()

    # ------------------------------------------------------------ page setup --
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.18)
        s.right_margin = Cm(3.18)

    def _setup_normal(self):
        st = self.doc.styles['Normal']
        st.font.size = PT_XIAOSI
        st.font.name = EN
        st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)
        pf = st.paragraph_format
        pf.line_spacing = LINE_EXACT_20
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    # ------------------------------------------------------------- primitives --
    def para(self, text='', cn_font=SONG, size=PT_XIAOSI, bold=False,
             en_font=EN, align=None, indent=None, hanging=None,
             space_before=None, space_after=None, color=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = LINE_EXACT_20
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        if align is not None:
            p.alignment = align
        if indent is not None:
            pf.first_line_indent = indent
        if hanging is not None:
            pf.left_indent = hanging
            pf.first_line_indent = -hanging
        if space_before is not None:
            pf.space_before = space_before
        if space_after is not None:
            pf.space_after = space_after
        if text:
            run = p.add_run(text)
            set_run_font(run, cn_font, size, bold, en_font, color)
        return p

    def heading(self, text, level=1, numbered=True, counters=None):
        """level 1..4, auto Chinese numbering."""
        counters = counters or {}
        prefix = ''
        if numbered:
            if level == 1:
                prefix = cn_number(counters.get(1, 1)) + '、'
            elif level == 2:
                prefix = '（' + cn_number(counters.get(2, 1)) + '）'
            elif level == 3:
                prefix = str(counters.get(3, 1)) + '.'
            elif level == 4:
                prefix = '（' + str(counters.get(4, 1)) + '）'
        full = prefix + text

        spec = {
            1: dict(size=Pt(16), align=WD_ALIGN_PARAGRAPH.CENTER,
                    sb=LINE_EXACT_20, sa=LINE_EXACT_20),
            2: dict(size=Pt(15), align=WD_ALIGN_PARAGRAPH.LEFT,
                    sb=None, sa=None),
            3: dict(size=Pt(14), align=WD_ALIGN_PARAGRAPH.LEFT,
                    sb=None, sa=None),
            4: dict(size=Pt(12), align=WD_ALIGN_PARAGRAPH.LEFT,
                    sb=None, sa=None),
        }[level]
        return self.para(full, cn_font=HEI, size=spec['size'], bold=False,
                         align=spec['align'], space_before=spec['sb'],
                         space_after=spec['sa'])

    def blank(self, n=1):
        for _ in range(n):
            self.para('')

    def page_break(self):
        self.doc.add_page_break()

    # ---------------------------------------------------------------- cover --
    def cover(self, data):
        for _ in range(5):
            self.para('')
        self.para(data.get('title', '论文题目'), cn_font=XBS, size=Pt(18),
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        if data.get('subtitle'):
            self.para('——' + data['subtitle'], cn_font=KAI, size=Pt(14),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        for _ in range(2):
            self.para('')
        if data.get('doc_type'):
            self.para(data['doc_type'], cn_font=HEI, size=Pt(16),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        for _ in range(3):
            self.para('')

        rows = [
            ('学  院', data.get('college', '')),
            ('专  业', data.get('major', '')),
            ('年  级', data.get('grade', '')),
            ('姓  名', data.get('author', '')),
            ('学  号', data.get('student_id', '')),
            ('指导教师', data.get('advisor', '')),
        ]
        table = self.doc.add_table(rows=len(rows), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (label, value) in enumerate(rows):
            c0, c1 = table.rows[i].cells
            c0.width = Cm(3.0)
            c1.width = Cm(8.0)
            for cell, txt, bold in ((c0, label, False), (c1, value or '____________', False)):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = cell.paragraphs[0].paragraph_format
                pf.line_spacing = LINE_EXACT_20
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                run = cell.paragraphs[0].add_run(txt)
                set_run_font(run, KAI, Pt(15), bold)
        self.blank(2)
        self.para(data.get('date', '____年__月'), cn_font=KAI, size=Pt(15),
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        self.page_break()

    # ------------------------------------------------------------- abstract --
    def abstract(self, data):
        self.para('摘要', cn_font=HEI, size=Pt(16),
                  align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_before=LINE_EXACT_20, space_after=LINE_EXACT_20)
        self.para(data['abstract'], indent=BODY_INDENT)
        self.blank(1)
        kw = data.get('keywords') or []
        if kw:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing = LINE_EXACT_20
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            r1 = p.add_run('关键词：')
            set_run_font(r1, SONG, PT_XIAOSI, True)
            r2 = p.add_run('，'.join(kw))
            set_run_font(r2, SONG, PT_XIAOSI, False)
        if data.get('abstract_en'):
            self.blank(1)
            self.para('Abstract', cn_font=EN, size=Pt(16), en_font=EN,
                      align=WD_ALIGN_PARAGRAPH.CENTER,
                      space_before=LINE_EXACT_20, space_after=LINE_EXACT_20)
            self.para(data['abstract_en'], en_font=EN)
            if data.get('keywords_en'):
                p = self.doc.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing = LINE_EXACT_20
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                r1 = p.add_run('Key words: ')
                set_run_font(r1, EN, PT_XIAOSI, True, EN)
                r2 = p.add_run(', '.join(data['keywords_en']))
                set_run_font(r2, EN, PT_XIAOSI, False, EN)
        self.page_break()

    # ----------------------------------------------------------------- toc --
    def toc(self):
        self.para('目  录', cn_font=HEI, size=Pt(16),
                  align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_before=LINE_EXACT_20, space_after=LINE_EXACT_20)
        self.para('（请在 Word 中插入自动目录：引用 → 目录 → 自动目录）',
                  size=PT_WU, color=(128, 128, 128))
        self.page_break()

    # --------------------------------------------------------------- blocks --
    def body_paragraph(self, text):
        self.para(text, indent=BODY_INDENT, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def table_block(self, caption, rows):
        if caption:
            self.para(caption, size=PT_WU, align=WD_ALIGN_PARAGRAPH.CENTER,
                      space_after=Pt(2))
        if not rows:
            return
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.rows[ri].cells[ci]
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = cell.paragraphs[0].paragraph_format
                pf.line_spacing = LINE_EXACT_20
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                run = cell.paragraphs[0].add_run(str(val))
                set_run_font(run, SONG, PT_WU, bold=(ri == 0))
        self.blank(1)

    def figure_block(self, caption, path):
        if path and Path(path).exists():
            try:
                self.doc.add_picture(str(path))
                self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:  # noqa: BLE001
                self.para('[图片插入失败: %s]' % path, size=PT_WU,
                          color=(128, 128, 128))
        if caption:
            self.para(caption, size=PT_WU, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(1)

    # ------------------------------------------------------------ chapters --
    def chapters(self, data):
        counters = {1: 0, 2: 0, 3: 0, 4: 0}
        for ch in data.get('chapters', []):
            numbered = ch.get('numbered', True)
            if numbered:
                counters[1] += 1
                counters[2] = counters[3] = counters[4] = 0
                self.heading(ch.get('title', ''), 1, True, counters)
            else:
                self.para(ch.get('title', ''), cn_font=HEI, size=Pt(16),
                          align=WD_ALIGN_PARAGRAPH.CENTER,
                          space_before=LINE_EXACT_20, space_after=LINE_EXACT_20)

            for sub in ch.get('subsections', []) or []:
                if numbered:
                    counters[2] += 1
                    counters[3] = counters[4] = 0
                    self.heading(sub.get('title', ''), 2, True, counters)
                else:
                    self.para(sub.get('title', ''), cn_font=HEI, size=Pt(15))
                self._emit(sub, 3, counters, numbered)

            if not ch.get('subsections'):
                self._emit(ch, 2, counters, numbered)

    def _emit(self, node, start_level, counters, numbered):
        """Render content string or blocks list."""
        blocks = node.get('blocks')
        if blocks:
            for b in blocks:
                t = b.get('type', 'p')
                if t == 'p':
                    self.body_paragraph(b.get('text', ''))
                elif t == 'table':
                    self.table_block(b.get('caption', ''), b.get('rows', []))
                elif t == 'figure':
                    self.figure_block(b.get('caption', ''), b.get('path'))
                elif t == 'h3':
                    if numbered:
                        counters[3] += 1
                        counters[4] = 0
                        self.heading(b.get('text', ''), 3, True, counters)
                    else:
                        self.para(b.get('text', ''), cn_font=HEI, size=Pt(14))
                elif t == 'h4':
                    if numbered:
                        counters[4] += 1
                        self.heading(b.get('text', ''), 4, True, counters)
                    else:
                        self.para(b.get('text', ''), cn_font=HEI, size=Pt(12))
            return
        content = node.get('content', '')
        for seg in content.split('\n\n'):
            seg = seg.strip()
            if seg:
                self.body_paragraph(seg)

    # ---------------------------------------------------------- references --
    def references(self, data):
        self.page_break()
        self.para('参考文献', cn_font=HEI, size=Pt(16),
                  align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_before=LINE_EXACT_20, space_after=LINE_EXACT_20)
        refs = data.get('references', []) or []
        if not refs:
            self.para('（参考文献由用户提供，不少于 15 篇）', size=PT_WU,
                      color=(128, 128, 128))
            return
        for i, ref in enumerate(refs, 1):
            text = ref if ref.strip().startswith('[') else '[%d] %s' % (i, ref)
            self.para(text, size=PT_WU, hanging=REF_HANGING,
                      align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # -------------------------------------------------------------- footer --
    def footer_page_number(self):
        footer = self.doc.sections[0].footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        set_run_font(run, EN, PT_WU, False, EN)
        fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
        run._element.append(fld1)
        run._element.append(instr)
        run._element.append(fld2)

    # ---------------------------------------------------------------- save --
    def save(self, path):
        self.doc.save(path)


def build(data, output_path):
    b = DocBuilder()
    b.cover(data)
    if data.get('abstract'):
        b.abstract(data)
    b.toc()
    b.chapters(data)
    b.references(data)
    b.footer_page_number()
    b.save(output_path)
    n_ch = len(data.get('chapters', []))
    n_ref = len(data.get('references', []) or [])
    print('开题报告已生成: %s' % output_path)
    print('  共 %d 个章节, %d 条参考文献' % (n_ch, n_ref))
    print('  待手工处理: 自动目录 / 脚注 / 罗马数字页码分节 / 原创性声明与授权声明页')


def main():
    ap = argparse.ArgumentParser(description='生成本科毕业论文开题报告 .docx')
    ap.add_argument('--output', required=True, help='输出 .docx 路径')
    ap.add_argument('--data', required=True, help='内容 JSON 文件路径')
    args = ap.parse_args()
    with open(args.data, 'r', encoding='utf-8') as f:
        data = json.load(f)
    build(data, args.output)


if __name__ == '__main__':
    main()
